#!/usr/bin/env python3
"""从资料库底稿生成两份文档：图片收录，以及对话与每日任务。"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "资料库" / "数据" / "library.json"

FONT = "PingFang SC"
INK = "1C1917"
NAVY = "1F3A5F"
MUTED = "44403C"
LINE = "D5DDE6"
HEAD_FILL = "E8EEF4"
BAND = "F4F7FB"
PAGE_W = 17.0


def load() -> dict:
    with DATA.open(encoding="utf-8") as fh:
        return json.load(fh)


def stamp(value: str) -> datetime:
    return datetime.strptime(value, "%Y-%m-%d %H:%M")


def set_run(run, text, size, bold=False, color=INK):
    run.text = text
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")
    return run


def add_text(paragraph, text, size=11, bold=False, color=INK):
    return set_run(paragraph.add_run(), text, size, bold, color)


def style_paragraph(paragraph, before=0, after=6, line=1.15, align="left"):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def valign(cell, value):
    tcpr = cell._tc.get_or_add_tcPr()
    node = tcpr.find(qn("w:vAlign"))
    if node is None:
        node = OxmlElement("w:vAlign")
        tcpr.append(node)
    node.set(qn("w:val"), value)


def cell_margins(cell, top=70, bottom=70, left=110, right=110):
    tcpr = cell._tc.get_or_add_tcPr()
    tcmar = tcpr.find(qn("w:tcMar"))
    if tcmar is None:
        tcmar = OxmlElement("w:tcMar")
        tcpr.append(tcmar)
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = tcmar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tcmar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tblpr = tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    total = int(sum(widths) * 567)
    tblw.set(qn("w:w"), str(total))
    tblw.set(qn("w:type"), "dxa")
    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), LINE)
    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tblpr.addnext(grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 567)))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(int(widths[idx] * 567)))
            tcw.set(qn("w:type"), "dxa")


def write_cell(cell, text, size=10.5, bold=False, color=INK, fill=None, align="left", center=True):
    if fill:
        shade(cell, fill)
    valign(cell, "center" if center else "top")
    cell_margins(cell)
    lines = text.split("\n") if isinstance(text, str) else list(text)
    for idx, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        paragraph.alignment = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
        }[align]
        style_paragraph(paragraph, 0, 1 if idx < len(lines) - 1 else 0, 1.08)
        if paragraph.runs:
            for run in list(paragraph.runs):
                run._element.getparent().remove(run._element)
        add_text(paragraph, line, size, bold, color)


def add_table(doc, headers, rows, widths, header=True, parent=None, spacer=True):
    host = doc if parent is None else parent
    table = host.add_table(rows=len(rows) + (1 if header else 0), cols=len(widths))
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_geometry(table, widths)
    offset = 0
    if header:
        for idx, label in enumerate(headers):
            write_cell(table.rows[0].cells[idx], label, size=10.5, bold=True, color=NAVY, fill=HEAD_FILL, center=True)
        trpr = table.rows[0]._tr.get_or_add_trPr()
        mark = OxmlElement("w:tblHeader")
        trpr.append(mark)
        offset = 1
    for ridx, row in enumerate(rows):
        for cidx, value in enumerate(row):
            write_cell(table.rows[ridx + offset].cells[cidx], value, center=False)
    for row in table.rows:
        trpr = row._tr.get_or_add_trPr()
        if trpr.find(qn("w:cantSplit")) is None:
            trpr.append(OxmlElement("w:cantSplit"))
    if spacer and parent is None:
        gap = doc.add_paragraph()
        style_paragraph(gap, 0, 2, 1.0)
    return table


def clear_borders(table):
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        node.set(qn("w:val"), "nil")
        node.set(qn("w:sz"), "0")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "auto")


def add_pdf_card(doc, note, placeholder):
    card = doc.add_table(rows=1, cols=1)
    card.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_geometry(card, [PAGE_W])
    clear_borders(card)
    row = card.rows[0]
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:cantSplit")) is None:
        trpr.append(OxmlElement("w:cantSplit"))
    cell = row.cells[0]
    cell_margins(cell, 20, 40, 0, 0)
    valign(cell, "top")
    title = cell.paragraphs[0]
    style_paragraph(title, 2, 1, 1.05)
    add_text(title, note["file"], 13, bold=True, color=NAVY)
    location = cell.add_paragraph()
    style_paragraph(location, 1, 0, 1.05)
    add_text(location, f"位置：{note['location']}", 10.5, color=MUTED)
    facts = cell.add_paragraph()
    style_paragraph(facts, 0, 3, 1.05)
    add_text(facts, f"大小：{note['size']}　　修改：{note['modified']}（北京时间）", 10.5, color=MUTED)
    add_table(
        doc,
        ["栏目", "内容"],
        [
            ["我要说的", blank_or(note["i_say"], placeholder)],
            ["我要做的", blank_or(note["i_do"], placeholder)],
            ["老板说的", blank_or(note["boss_said"], placeholder)],
        ],
        [3.2, 13.8],
        parent=cell,
        spacer=False,
    )
    gap = doc.add_paragraph()
    style_paragraph(gap, 2, 2, 1.0)


def merge_note(table, text):
    row = table.rows[-1]
    anchor = row.cells[0]
    for cell in row.cells[1:]:
        anchor.merge(cell)
    write_cell(anchor, text, color=MUTED, fill="FBFCFD", center=False)


def paragraph_border(paragraph, edge, color, sz="12", space="1"):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    node = OxmlElement(f"w:{edge}")
    node.set(qn("w:val"), "single")
    node.set(qn("w:sz"), sz)
    node.set(qn("w:space"), space)
    node.set(qn("w:color"), color)
    pbdr.append(node)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    fmt = normal.paragraph_format
    fmt.space_after = Pt(6)
    fmt.line_spacing = 1.15
    for style_name, size, before, after in (
        ("Heading 1", 16, 12, 4),
        ("Heading 2", 13, 8, 3),
    ):
        style = doc.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(NAVY)
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), FONT)
        pf = style.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = 1.1
        pf.keep_with_next = True


def apply_section(doc, header_label, updated_at):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.15)
    section.bottom_margin = Cm(1.8)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.5)
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    style_paragraph(hp, 0, 2, 1.0)
    add_text(hp, header_label, 9, color=MUTED)
    paragraph_border(hp, "bottom", LINE, "8", "1")
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    style_paragraph(fp, 2, 0, 1.0)
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(PAGE_W), WD_TAB_ALIGNMENT.RIGHT)
    add_text(fp, f"更新于 {updated_at}（北京时间）", 9, color=MUTED)
    add_text(fp, "\t第 ", 9, color=MUTED)
    run = fp.add_run()
    set_run(run, "", 9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    add_text(fp, " 页", 9, color=MUTED)
    paragraph_border(fp, "top", LINE, "8", "1")
    settings = doc.settings.element
    lang = OxmlElement("w:themeFontLang")
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")
    settings.append(lang)


def new_doc(title, subject, updated_at):
    doc = Document()
    configure_styles(doc)
    props = doc.core_properties
    props.title = title
    props.subject = subject
    props.author = "资料库"
    props.last_modified_by = "资料库"
    props.category = "资料库"
    props.created = stamp(updated_at)
    props.modified = stamp(updated_at)
    return doc


def add_title(doc, title, subtitle):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, 0, 2, 1.05)
    add_text(paragraph, title, 22, bold=True, color=NAVY)
    sub = doc.add_paragraph()
    style_paragraph(sub, 0, 8, 1.05)
    add_text(sub, subtitle, 11, color=MUTED)
    paragraph_border(sub, "bottom", NAVY, "12", "2")


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        set_run(run, run.text, 16 if level == 1 else 13, bold=True, color=NAVY)
    return paragraph


def add_body(doc, text, size=11, color=INK, after=6, before=0):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, before, after, 1.15)
    add_text(paragraph, text, size, color=color)
    return paragraph


def add_rule_box(doc, title, rules):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_geometry(table, [PAGE_W])
    cell = table.cell(0, 0)
    shade(cell, BAND)
    valign(cell, "top")
    cell_margins(cell, 90, 80, 140, 140)
    first = cell.paragraphs[0]
    style_paragraph(first, 0, 4, 1.08)
    add_text(first, title, 11, bold=True, color=NAVY)
    for idx, rule in enumerate(rules, start=1):
        paragraph = cell.add_paragraph()
        style_paragraph(paragraph, 1, 1, 1.08)
        add_text(paragraph, f"{idx}.  {rule}", 10.5, color=INK)
    spacer = doc.add_paragraph()
    style_paragraph(spacer, 2, 0, 1.0)


def blank_or(value, placeholder):
    text = (value or "").strip()
    return text if text else placeholder


def build_image_doc(data, record):
    updated = data["meta"]["updated_at"]
    doc = new_doc("图片收录", record["id"], updated)
    apply_section(doc, "资料库  ·  图片收录", updated)
    add_title(doc, "图片收录", f"{record['id']}    {record['recorded_at']}（北京时间）")
    add_body(doc, record["scope"], color=MUTED, after=8)
    rows = [
        ["收录编号", record["id"]],
        ["收录时间", f"{record['recorded_at']}（北京时间）"],
        ["来源", record["source"]],
        ["访达窗口", record["finder_title"]],
        ["路径栏", "  ›  ".join(record["breadcrumb"])],
        ["本机路径", record["posix_path"]],
        ["视图", record["view"]],
        ["原图", record["image"]],
        ["原图尺寸", f"{record['image_pixels']}，{record['image_size']}"],
        ["本说明", record["doc"]],
    ]
    add_heading(doc, "收录信息", 1)
    add_table(doc, ["项目", "内容"], rows, [3.4, 13.6])
    add_heading(doc, "原图", 1)
    image_path = ROOT / record["image"]
    if not image_path.is_file():
        raise FileNotFoundError(image_path)
    pic = doc.add_paragraph()
    style_paragraph(pic, 2, 4, 1.0)
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pic.add_run()
    run.add_picture(str(image_path), width=Cm(16.2))
    add_heading(doc, "图中的文件", 1)
    file_rows = []
    for item in record["files"]:
        modified = item["modified"].replace(" ", "\n", 1)
        file_rows.append([item["name"], modified, item["modified_label"], item["size"], item["kind"]])
    add_table(
        doc,
        ["名称", "修改时间", "访达原文", "大小", "种类"],
        file_rows,
        [7.0, 3.0, 2.5, 1.9, 2.6],
    )
    add_body(doc, record["date_basis"], size=10.5, color=MUTED, before=0, after=6)
    add_heading(doc, "画面上的标记", 1)
    for mark in record["marks"]:
        add_body(doc, mark, size=11, after=4)
    add_heading(doc, "和主文档的关系", 1)
    add_body(
        doc,
        "对话、关于 PDF 的三栏、老板要求、每日任务写在「资料库/对话与每日任务/工程采购学习任务.docx」。这份文档只负责这张图。",
    )
    out = ROOT / record["doc"]
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def build_main_doc(data):
    meta = data["meta"]
    updated = meta["updated_at"]
    doc = new_doc(meta["title"], "对话记录与每日计划", updated)
    apply_section(doc, "资料库  ·  对话与每日任务", updated)
    add_title(doc, meta["title"], f"对话记录与每日计划    更新于 {updated}（北京时间）")
    add_body(doc, meta["intro"])
    add_rule_box(doc, meta["rules_title"], meta["rules"])

    add_heading(doc, "一、对话记录", 1)
    add_body(doc, "下表只留重要对话，以及你点名要记的对话。", size=10.5, color=MUTED)
    conv_rows = []
    for item in data["conversations"]:
        lines = [f"{idx}. {line}" for idx, line in enumerate(item["summary"], start=1)]
        lines.append("记入：" + "、".join(item["filed_under"]))
        conv_rows.append([item["time"], item["record_reason"], "\n".join(lines)])
    add_table(doc, ["时间", "为何记下", "内容"], conv_rows, [3.5, 4.3, 9.2])

    add_heading(doc, "二、关于这些 PDF", 1)
    add_body(doc, meta["pdf_intro"])
    placeholder = meta["empty_note"]
    for note in data["pdf_notes"]:
        add_pdf_card(doc, note, placeholder)

    add_heading(doc, "三、部分一：老板要求", 1)
    add_body(doc, meta["boss_intro"])
    headers = ["日期", "教学录音", "老板要求", "对应", "状态"]
    widths = [2.5, 3.5, 4.4, 4.2, 2.4]
    if data["boss_requirements"]:
        rows = [
            [
                item["date"],
                item["recording"],
                item["requirement"],
                item["corresponds_to"],
                item["status"],
            ]
            for item in data["boss_requirements"]
        ]
        add_table(doc, headers, rows, widths)
    else:
        table = add_table(doc, headers, [["", "", "", "", ""]], widths)
        merge_note(table, meta["empty_boss"])

    add_heading(doc, "四、部分二：每日任务和计划", 1)
    add_body(doc, meta["tasks_intro"])
    for day in data["daily_tasks"]:
        add_heading(doc, f"{day['date']}　{day['weekday']}", 2)
        add_body(doc, day["date_source"], size=10.5, color=MUTED, after=4)
        rows = [[task["text"], task["status"], task["note"]] for task in day["tasks"]]
        add_table(doc, ["任务", "状态", "备注"], rows, [9.4, 2.2, 5.4])

    add_heading(doc, "五、资料放在哪里", 1)
    folders = meta["folders"]
    place_rows = []
    for record in data["image_records"]:
        place_rows.append([record["folder"], f"原图 {record['id']}"])
        place_rows.append([record["doc"], "这张图的说明"])
    place_rows.extend(
        [
            [folders["audio"], "教学录音"],
            [data["outputs"]["main_doc"], "本主文档"],
            [folders["data"], "底稿，改这里再生成"],
        ]
    )
    add_table(doc, ["位置", "放什么"], place_rows, [10.6, 6.4])
    out = ROOT / data["outputs"]["main_doc"]
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def main():
    data = load()
    written = [build_main_doc(data)]
    written.extend(build_image_doc(data, record) for record in data["image_records"])
    for path in written:
        print(path.relative_to(ROOT))


if __name__ == "__main__":
    main()
